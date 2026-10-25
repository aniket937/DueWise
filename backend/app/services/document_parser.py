import re
import io
from pathlib import Path
from typing import Tuple, Optional
import pymupdf  # PyMuPDF
from docx import Document as DocxDocument
from app.services.ocr_service import is_tesseract_installed, extract_text_from_image

MAX_WORD_LIMIT = 50000

def sanitize_filename(name: str) -> str:
    """Sanitize filename to prevent directory traversal and remove dangerous chars."""
    name = Path(name).name
    # Keep alphanumeric, dots, underscores, dashes
    clean = re.sub(r"[^\w\.\- ]", "_", name)
    return clean[:200] if clean else "unnamed_document"

def extract_from_txt(file_bytes: bytes) -> str:
    """Extract text from TXT with multiple encoding fallbacks."""
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")

def extract_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX document."""
    doc_io = io.BytesIO(file_bytes)
    doc = DocxDocument(doc_io)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)

def extract_from_pdf(file_bytes: bytes) -> Tuple[str, int]:
    """Extract text from PDF using PyMuPDF, with graceful OCR fallback for scanned pages."""
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    total_pages = len(doc)
    text_parts = []

    for page_idx in range(total_pages):
        page = doc.load_page(page_idx)
        text = page.get_text("text").strip()

        # If page has virtually no text, check if scanned page can be OCR'd
        if not text and is_tesseract_installed():
            try:
                pix = page.get_pixmap()
                from PIL import Image
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_text = extract_text_from_image(img)
                if ocr_text:
                    text = ocr_text.strip()
            except Exception:
                pass

        if text:
            text_parts.append(f"--- Page {page_idx + 1} ---\n{text}")

    doc.close()
    full_text = "\n\n".join(text_parts)
    return full_text, total_pages

def parse_document(file_bytes: bytes, filename: str, mime_type: str) -> Tuple[str, Optional[int], int]:
    """
    Parse uploaded document and return (extracted_text, page_count, word_count).
    Ensures safe bounds and clean extraction.
    """
    ext = Path(filename).suffix.lower()
    page_count: Optional[int] = None
    extracted_text = ""

    if ext == ".pdf" or "pdf" in mime_type:
        extracted_text, page_count = extract_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"] or "wordprocessingml" in mime_type:
        extracted_text = extract_from_docx(file_bytes)
        page_count = 1
    elif ext in [".txt", ".text", ".md"] or "text/plain" in mime_type:
        extracted_text = extract_from_txt(file_bytes)
        page_count = 1
    else:
        raise ValueError(f"Unsupported file format: {ext or mime_type}. Supported formats: PDF, DOCX, TXT.")

    # Clean whitespace
    extracted_text = re.sub(r"\r\n|\r", "\n", extracted_text)
    extracted_text = re.sub(r"\n{3,}", "\n\n", extracted_text).strip()

    words = extracted_text.split()
    word_count = len(words)

    # Bound word count if excessively large
    if word_count > MAX_WORD_LIMIT:
        extracted_text = " ".join(words[:MAX_WORD_LIMIT]) + "\n\n[Content truncated to maximum processing limit]"
        word_count = MAX_WORD_LIMIT

    return extracted_text, page_count, word_count
